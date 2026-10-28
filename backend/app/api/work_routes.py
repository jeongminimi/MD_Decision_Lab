"""Operational workspace: observed sales stay separate from MD-entered assumptions."""

import csv
import io
import math
from collections import defaultdict
from datetime import date, timedelta

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from pydantic import BaseModel, Field
from typing import Literal
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..database import get_db
from ..models import WorkDailySale, WorkDataset, WorkDecision, WorkProduct, WorkProductInput, WorkInputSnapshot
from ..services.work_risk import RISK_PRIORITY, classify_work_risk

router = APIRouter(prefix="/api/work", tags=["MD workspaces"])
MAX_UPLOAD = 15 * 1024 * 1024


def _number(value: object, label: str, *, integer: bool = False) -> int | float | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        result = int(str(value).strip()) if integer else float(str(value).strip())
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{label}: 숫자를 입력하세요.") from exc
    if not math.isfinite(result) or result < 0:
        raise ValueError(f"{label}: 0 이상의 유효한 숫자를 입력하세요.")
    return result


async def _rows(file: UploadFile) -> list[dict[str, str]]:
    if not file.filename:
        raise ValueError("파일 이름이 없습니다.")
    raw = await file.read(MAX_UPLOAD + 1)
    if len(raw) > MAX_UPLOAD:
        raise ValueError("파일 하나당 15MB 이하의 CSV·엑셀 파일을 올려주세요.")
    name = file.filename.lower()
    if name.endswith(".csv"):
        try:
            content = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            content = raw.decode("cp949")
        reader = csv.DictReader(io.StringIO(content))
        if not reader.fieldnames:
            raise ValueError("첫 줄에 열 이름이 필요합니다.")
        return [{str(k).strip(): (v or "").strip() for k, v in row.items() if k is not None} for row in reader]
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        try:
            sheet = book.active
            iterator = sheet.iter_rows(values_only=True)
            headers = [str(x).strip() if x is not None else "" for x in next(iterator)]
            return [{key: "" if value is None else str(value).strip() for key, value in zip(headers, row) if key}
                    for row in iterator if any(value is not None for value in row)]
        finally:
            book.close()
    raise ValueError("CSV 또는 XLSX 파일을 올려주세요.")


def _require_fields(rows: list[dict[str, str]], required: set[str], label: str) -> None:
    if not rows:
        raise ValueError(f"{label} 파일이 비어 있습니다.")
    missing = required - rows[0].keys()
    if missing:
        raise ValueError(f"{label} 필수 열이 없습니다: {', '.join(sorted(missing))}")


@router.get("/datasets")
def datasets(db: Session = Depends(get_db)):
    entries = db.scalars(select(WorkDataset).order_by(WorkDataset.created_at.desc(), WorkDataset.id.desc())).all()
    return [{"id": d.id, "name": d.name, "kind": d.kind, "start_date": d.start_date,
             "end_date": d.end_date, "source_note": d.source_note,
             "complete_calendar": d.complete_calendar} for d in entries]


@router.post("/datasets/upload")
async def upload_dataset(
    name: str = Form(...), snapshot_date: date = Form(...), complete_calendar: bool = Form(False),
    products_file: UploadFile = File(...), sales_file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    """Create a new company dataset atomically; never mutate another workspace."""
    if not name.strip() or len(name) > 120:
        raise HTTPException(400, "작업공간 이름은 1~120자로 입력하세요.")
    try:
        product_rows, sale_rows = await _rows(products_file), await _rows(sales_file)
        _require_fields(product_rows, {"sku", "name"}, "상품")
        _require_fields(sale_rows, {"sku", "sale_date", "units_sold"}, "판매")
        products: dict[str, dict] = {}
        for index, row in enumerate(product_rows, 2):
            sku = row["sku"].strip()
            if not sku or len(sku) > 64 or sku in products or not row["name"].strip():
                raise ValueError(f"상품 {index}행: SKU 중복·누락 또는 상품명 누락입니다.")
            fields = {key: _number(row.get(key), f"상품 {index}행 {key}", integer=key in {
                "stock_on_hand", "inbound_units", "lead_time_days"}) for key in (
                "stock_on_hand", "inbound_units", "lead_time_days", "unit_price_krw", "unit_cost_krw")}
            reorder = row.get("reorderable", "").strip().lower()
            if reorder not in {"", "true", "false", "1", "0", "yes", "no"}:
                raise ValueError(f"상품 {index}행 reorderable: true 또는 false를 입력하세요.")
            fields["reorderable"] = None if not reorder else reorder in {"true", "1", "yes"}
            if row.get("lifecycle", "") not in {"", "regular", "summer_clearance", "fall_prelaunch", "winter_carryover"}:
                raise ValueError(f"상품 {index}행 lifecycle: 허용되지 않는 상품 상태입니다.")
            products[sku] = {"row": row, "inputs": fields}
        sales: dict[tuple[str, date, str], dict] = {}
        for index, row in enumerate(sale_rows, 2):
            sku = row["sku"].strip()
            if sku not in products:
                raise ValueError(f"판매 {index}행: 상품 파일에 없는 SKU {sku}")
            day = date.fromisoformat(row["sale_date"].strip()[:10])
            if day > snapshot_date:
                raise ValueError(f"판매 {index}행: 재고 기준일보다 미래의 판매일입니다.")
            units = _number(row.get("units_sold"), f"판매 {index}행 units_sold", integer=True)
            if units is None:
                raise ValueError(f"판매 {index}행: 판매 수량이 비어 있습니다.")
            revenue = _number(row.get("revenue"), f"판매 {index}행 revenue")
            profit = row.get("gross_margin", "")
            margin = None if profit == "" else float(profit)
            if margin is not None and not math.isfinite(margin):
                raise ValueError(f"판매 {index}행 gross_margin: 유효한 숫자가 아닙니다.")
            if (revenue is None and margin is not None) or (units == 0 and revenue not in {None, 0}):
                raise ValueError(f"판매 {index}행: 판매·매출·이익 값이 서로 맞지 않습니다.")
            channel = (row.get("channel") or "전체").strip()[:32]
            key = (sku, day, channel)
            if key not in sales:
                sales[key] = {"units": 0, "revenue": 0.0 if revenue is not None else None,
                              "gross_profit": 0.0 if margin is not None else None}
            item = sales[key]
            item["units"] += units
            item["revenue"] = None if revenue is None or item["revenue"] is None else item["revenue"] + revenue
            item["gross_profit"] = None if margin is None or item["gross_profit"] is None else item["gross_profit"] + margin
        first, last = min(key[1] for key in sales), max(key[1] for key in sales)
        if snapshot_date != last and any(p["inputs"]["stock_on_hand"] is not None for p in products.values()):
            raise ValueError("현재고를 올릴 때 재고 기준일은 마지막 판매일과 같아야 합니다.")
        dataset = WorkDataset(name=name.strip(), kind="company", start_date=first, end_date=last,
                              complete_calendar=complete_calendar,
                              source_note="자사 파일 업로드 · 판매는 관측값, 재고·가격·원가는 업로드 기준일 입력값")
        db.add(dataset)
        db.flush()
        ids = {}
        for sku, entry in products.items():
            row = entry["row"]
            product = WorkProduct(dataset_id=dataset.id, sku=sku, name=row["name"].strip()[:180],
                                  category=(row.get("category") or "기타")[:100],
                                  section=(row.get("section") or "전체")[:100],
                                  lifecycle=(row.get("lifecycle") or None))
            db.add(product)
            db.flush()
            ids[sku] = product.id
            inputs = entry["inputs"]
            if any(value is not None for value in inputs.values()):
                db.add(WorkProductInput(product_id=product.id, as_of_date=snapshot_date,
                                        source_label="company_upload", **inputs))
                db.add(WorkInputSnapshot(product_id=product.id, as_of_date=snapshot_date,
                                         source_label="company_upload", **inputs))
        for (sku, day, channel), values in sales.items():
            db.add(WorkDailySale(product_id=ids[sku], sale_date=day, channel=channel, **values))
        db.commit()
        return {"id": dataset.id, "product_count": len(products), "daily_rows": len(sales),
                "start_date": first, "end_date": last}
    except (ValueError, KeyError, UnicodeError, StopIteration, OSError) as exc:
        db.rollback()
        raise HTTPException(400, f"파일을 반영하지 않았습니다: {exc}") from exc


@router.post("/datasets/{dataset_id}/sales/import")
async def append_company_sales(dataset_id: int, complete_calendar: bool = Form(False),
                               sales_file: UploadFile = File(...), db: Session = Depends(get_db)):
    """Upsert dated sales without deleting unrelated rows, so post-action periods can arrive later."""
    dataset = _dataset(dataset_id, db)
    if dataset.kind != "company":
        raise HTTPException(400, "H&M 원본 판매 기록은 수정할 수 없습니다.")
    try:
        rows = await _rows(sales_file)
        _require_fields(rows, {"sku", "sale_date", "units_sold"}, "판매")
        products = {p.sku: p.id for p in db.scalars(select(WorkProduct).where(
            WorkProduct.dataset_id == dataset_id)).all()}
        parsed: dict[tuple[int, date, str], dict] = {}
        for index, row in enumerate(rows, 2):
            sku = row["sku"].strip()
            if sku not in products:
                raise ValueError(f"{index}행: 등록되지 않은 SKU {sku}")
            day = date.fromisoformat(row["sale_date"].strip()[:10])
            units = _number(row.get("units_sold"), f"{index}행 units_sold", integer=True)
            if units is None:
                raise ValueError(f"{index}행: 판매 수량이 비어 있습니다.")
            revenue = _number(row.get("revenue"), f"{index}행 revenue")
            margin_raw = row.get("gross_margin", "")
            margin = None if margin_raw == "" else float(margin_raw)
            if margin is not None and (not math.isfinite(margin) or revenue is None):
                raise ValueError(f"{index}행: 매출총이익과 매출 값이 맞지 않습니다.")
            if units == 0 and revenue not in {None, 0}:
                raise ValueError(f"{index}행: 판매 0개일 때 매출은 0이어야 합니다.")
            key = (products[sku], day, (row.get("channel") or "전체").strip()[:32])
            if key not in parsed:
                parsed[key] = {"units": 0, "revenue": 0.0 if revenue is not None else None,
                               "gross_profit": 0.0 if margin is not None else None}
            value = parsed[key]
            value["units"] += units
            value["revenue"] = None if revenue is None or value["revenue"] is None else value["revenue"] + revenue
            value["gross_profit"] = None if margin is None or value["gross_profit"] is None else value["gross_profit"] + margin
        for (product_id, day, channel), values in parsed.items():
            existing = db.scalar(select(WorkDailySale).where(WorkDailySale.product_id == product_id,
                WorkDailySale.sale_date == day, WorkDailySale.channel == channel))
            if existing is None:
                db.add(WorkDailySale(product_id=product_id, sale_date=day, channel=channel, **values))
            else:
                for key, value in values.items():
                    setattr(existing, key, value)
        dataset.start_date = min(dataset.start_date, min(key[1] for key in parsed))
        dataset.end_date = max(dataset.end_date, max(key[1] for key in parsed))
        dataset.complete_calendar = dataset.complete_calendar and complete_calendar
        db.commit()
        return {"updated": len(parsed), "end_date": dataset.end_date,
                "complete_calendar": dataset.complete_calendar}
    except (ValueError, KeyError, UnicodeError, StopIteration, OSError) as exc:
        db.rollback()
        raise HTTPException(400, f"판매 기록을 반영하지 않았습니다: {exc}") from exc


def _dataset(dataset_id: int, db: Session) -> WorkDataset:
    value = db.get(WorkDataset, dataset_id)
    if value is None:
        raise HTTPException(404, "작업공간을 찾지 못했습니다.")
    return value


def _product(dataset_id: int, sku: str, db: Session) -> WorkProduct:
    value = db.scalar(select(WorkProduct).where(WorkProduct.dataset_id == dataset_id, WorkProduct.sku == sku))
    if value is None:
        raise HTTPException(404, "해당 상품을 찾지 못했습니다.")
    return value


def _input_on(product_id: int, day: date, db: Session):
    snapshot = db.scalar(select(WorkInputSnapshot).where(
        WorkInputSnapshot.product_id == product_id, WorkInputSnapshot.as_of_date == day))
    if snapshot is not None:
        return snapshot
    legacy = db.get(WorkProductInput, product_id)
    return legacy if legacy is not None and legacy.as_of_date == day else None


def _previous_year(day: date) -> date:
    try:
        return day.replace(year=day.year - 1)
    except ValueError:
        return day.replace(year=day.year - 1, day=28)


def _units(rows: list[WorkDailySale], start: date, end: date) -> int:
    return sum(row.units for row in rows if start <= row.sale_date <= end)


@router.get("/analysis")
def analysis(dataset_id: int, as_of: date | None = None, section: str | None = None,
             chart_start: date | None = None, db: Session = Depends(get_db)):
    dataset = _dataset(dataset_id, db)
    day = as_of or dataset.end_date
    if not dataset.start_date <= day <= dataset.end_date:
        raise HTTPException(400, "분석 기준일은 데이터에 포함된 날짜여야 합니다.")
    start = chart_start or max(dataset.start_date, day - timedelta(days=55))
    if not dataset.start_date <= start <= day:
        raise HTTPException(400, "그래프 시작일은 판매 기간 안에서 기준일보다 빠르게 선택하세요.")
    query = select(WorkProduct).where(WorkProduct.dataset_id == dataset_id)
    if section:
        query = query.where(WorkProduct.section == section)
    products = db.scalars(query.order_by(WorkProduct.sku)).all()
    if not products:
        return {"dataset": {"id": dataset.id, "name": dataset.name, "kind": dataset.kind,
                            "source_note": dataset.source_note}, "as_of_date": day, "chart_start": start,
                "sections": [], "products": [], "daily_chart": [], "summary": {"total_products": 0,
                "attention_count": 0, "units_30d": 0}}
    ids = [p.id for p in products]
    min_day = min(start, day - timedelta(days=59), _previous_year(day) - timedelta(days=29))
    sale_rows = db.scalars(select(WorkDailySale).where(WorkDailySale.product_id.in_(ids),
                        WorkDailySale.sale_date.between(min_day, day))).all()
    by_product: dict[int, list[WorkDailySale]] = defaultdict(list)
    chart: dict[date, int] = defaultdict(int)
    for row in sale_rows:
        by_product[row.product_id].append(row)
        if start <= row.sale_date <= day:
            chart[row.sale_date] += row.units
    inputs = {value.product_id: value for value in db.scalars(
        select(WorkInputSnapshot).where(WorkInputSnapshot.product_id.in_(ids),
                                        WorkInputSnapshot.as_of_date == day)).all()}
    legacy_inputs = {value.product_id: value for value in db.scalars(
        select(WorkProductInput).where(WorkProductInput.product_id.in_(ids),
                                       WorkProductInput.as_of_date == day)).all()}
    previous_day = _previous_year(day)
    sold_before = set(db.scalars(select(WorkDailySale.product_id).where(
        WorkDailySale.product_id.in_(ids), WorkDailySale.sale_date <= day, WorkDailySale.units > 0).distinct()))
    items = []
    for product in products:
        rows = by_product[product.id]
        u7 = _units(rows, day - timedelta(days=6), day)
        u14 = _units(rows, day - timedelta(days=13), day)
        p14 = _units(rows, day - timedelta(days=27), day - timedelta(days=14))
        u30 = _units(rows, day - timedelta(days=29), day)
        p30 = _units(rows, day - timedelta(days=59), day - timedelta(days=30))
        y30 = _units(rows, previous_day - timedelta(days=29), previous_day)
        channel_totals: dict[str, int] = defaultdict(int)
        for sale in rows:
            if day - timedelta(days=29) <= sale.sale_date <= day:
                channel_totals[sale.channel] += sale.units
        has_year = dataset.start_date <= previous_day - timedelta(days=29)
        data_days = min(30, (day - dataset.start_date).days + 1)
        inp = inputs.get(product.id) or legacy_inputs.get(product.id)
        valid_input = inp is not None and inp.as_of_date == day
        stock = inp.stock_on_hand if valid_input else None
        inbound = inp.inbound_units if valid_input else None
        price = inp.unit_price_krw if valid_input else None
        cost = inp.unit_cost_krw if valid_input else None
        rate = u14 / 14 if data_days >= 28 else u30 / data_days
        cover = round(stock / rate, 1) if stock is not None and rate > 0 else None
        if data_days < 28 or not dataset.complete_calendar:
            risk, reason, action = "data_gap", "비교 가능한 판매 기간 또는 기록 완전성 확인이 필요합니다.", "판매 자료 확인"
        else:
            risk, reason, action = classify_work_risk(
                u14=u14, p14=p14, stock=stock, cover=cover,
                lead_time_days=inp.lead_time_days if valid_input else None,
                reorderable=inp.reorderable if valid_input else None,
                observed_any=product.id in sold_before)
        period_rows = [row for row in rows if day - timedelta(days=29) <= row.sale_date <= day]
        observed_revenue = sum(row.revenue or 0 for row in period_rows)
        revenue_available = dataset.kind == "company" and bool(period_rows) and all(row.revenue is not None for row in period_rows)
        observed_profit = sum(row.gross_profit or 0 for row in period_rows)
        profit_available = dataset.kind == "company" and bool(period_rows) and all(row.gross_profit is not None for row in period_rows)
        items.append({"sku": product.sku, "name": product.name, "category": product.category,
                      "section": product.section, "image_ref": product.image_ref,
                      "lifecycle": product.lifecycle, "units_7d": u7, "units_14d": u14,
                      "prior_14d": p14, "units_30d": u30, "prior_30d": p30,
                      "year_ago_30d": y30 if has_year else None,
                      "channel_units_30d": dict(sorted(channel_totals.items())),
                      "observed_revenue_30d": round(observed_revenue, 2) if revenue_available else None,
                      "observed_gross_profit_30d": round(observed_profit, 2) if profit_available else None,
                      "stock_on_hand": stock, "inbound_units": inbound,
                      "lead_time_days": inp.lead_time_days if valid_input else None,
                      "unit_price_krw": price, "unit_cost_krw": cost,
                      "reorderable": inp.reorderable if valid_input else None,
                      "input_source": inp.source_label if valid_input else None,
                      "input_as_of_date": inp.as_of_date if inp else None,
                      "days_of_cover": cover, "risk": risk, "reason": reason, "action": action})
    items.sort(key=lambda item: (RISK_PRIORITY[item["risk"]], -item["units_30d"], item["sku"]))
    sections = db.scalars(select(WorkProduct.section).where(WorkProduct.dataset_id == dataset_id).distinct().order_by(WorkProduct.section)).all()
    daily_chart = [{"date": (start + timedelta(days=i)).isoformat(), "units": chart[start + timedelta(days=i)]}
                   for i in range((day - start).days + 1)]
    return {"dataset": {"id": dataset.id, "name": dataset.name, "kind": dataset.kind,
                        "source_note": dataset.source_note, "complete_calendar": dataset.complete_calendar},
            "as_of_date": day, "chart_start": start, "sections": sections,
            "products": items, "daily_chart": daily_chart,
            "summary": {"total_products": len(items),
                        "attention_count": sum(item["risk"] not in {"stable", "data_gap", "not_observed"} for item in items),
                        "units_30d": sum(item["units_30d"] for item in items)}}


class WorkInputPayload(BaseModel):
    as_of_date: date
    stock_on_hand: int | None = Field(default=None, ge=0)
    inbound_units: int | None = Field(default=None, ge=0)
    lead_time_days: int | None = Field(default=None, ge=0)
    unit_price_krw: float | None = Field(default=None, ge=0, allow_inf_nan=False)
    unit_cost_krw: float | None = Field(default=None, ge=0, allow_inf_nan=False)
    reorderable: bool | None = None
    lifecycle: Literal["regular", "summer_clearance", "fall_prelaunch", "winter_carryover"] | None = None


@router.put("/datasets/{dataset_id}/products/{sku}/inputs")
def update_inputs(dataset_id: int, sku: str, payload: WorkInputPayload, db: Session = Depends(get_db)):
    dataset = _dataset(dataset_id, db)
    if not dataset.start_date <= payload.as_of_date <= dataset.end_date:
        raise HTTPException(400, "입력 기준일은 판매 데이터 기간 안에서 선택하세요.")
    product = _product(dataset_id, sku, db)
    inp = db.get(WorkProductInput, product.id)
    if inp is None:
        inp = WorkProductInput(product_id=product.id)
        db.add(inp)
    for key, value in payload.model_dump(exclude={"lifecycle"}).items():
        setattr(inp, key, value)
    product.lifecycle = payload.lifecycle
    inp.source_label = "simulation" if dataset.kind == "hm" else "md_input"
    snapshot = db.scalar(select(WorkInputSnapshot).where(
        WorkInputSnapshot.product_id == product.id,
        WorkInputSnapshot.as_of_date == payload.as_of_date))
    if snapshot is None:
        snapshot = WorkInputSnapshot(product_id=product.id, as_of_date=payload.as_of_date)
        db.add(snapshot)
    for key, value in payload.model_dump(exclude={"lifecycle"}).items():
        setattr(snapshot, key, value)
    snapshot.source_label = inp.source_label
    db.commit()
    return {"sku": sku, "source_label": inp.source_label, "as_of_date": inp.as_of_date}


class WorkComparePayload(BaseModel):
    dataset_id: int
    sku: str
    as_of_date: date
    discount_pct: float = Field(ge=0, le=70, allow_inf_nan=False)
    expected_units_30d: int = Field(ge=0, le=1_000_000)
    purchase_units: int = Field(default=0, ge=0, le=1_000_000)


def _compare(payload: WorkComparePayload, db: Session) -> dict:
    dataset = _dataset(payload.dataset_id, db)
    product = _product(payload.dataset_id, payload.sku, db)
    inp = _input_on(product.id, payload.as_of_date, db)
    if inp is None or any(getattr(inp, key) is None for key in
        ("stock_on_hand", "inbound_units", "unit_price_krw", "unit_cost_krw")):
        raise HTTPException(409, "같은 기준일의 현재고·입고 예정·원화 판매가·원가를 먼저 입력하세요. 입고가 없으면 0을 입력하세요.")
    if payload.purchase_units and inp.reorderable is not True:
        raise HTTPException(400, "재발주 가능 여부가 확인된 상품만 추가 매입을 가정할 수 있습니다.")
    if payload.as_of_date > dataset.end_date:
        raise HTTPException(400, "판매 기록 이후의 기준일은 비교할 수 없습니다.")
    available = inp.stock_on_hand + (inp.inbound_units or 0) + payload.purchase_units
    price = round(inp.unit_price_krw * (1 - payload.discount_pct / 100), 2)
    feasible = payload.expected_units_30d <= available
    profit = round(payload.expected_units_30d * (price - inp.unit_cost_krw), 2) if feasible else None
    return {"sku": product.sku, "dataset_kind": dataset.kind, "as_of_date": payload.as_of_date,
            "source": "H&M 관측 판매 + 사용자 시뮬레이션" if dataset.kind == "hm" else "자사 판매 + MD 입력 시나리오",
            "discount_pct": payload.discount_pct, "expected_units_30d": payload.expected_units_30d,
            "purchase_units": payload.purchase_units, "available_units": available,
            "proposed_unit_price_krw": price, "feasible": feasible,
            "estimated_revenue_krw": round(payload.expected_units_30d * price, 2) if feasible else None,
            "estimated_gross_profit_krw": profit,
            "estimated_remaining_units": available - payload.expected_units_30d if feasible else None,
            "note": "입력한 예상 판매량을 가정한 계산입니다. 할인 효과나 실제 매출 증가를 검증한 결과가 아닙니다."}


@router.post("/compare")
def compare(payload: WorkComparePayload, db: Session = Depends(get_db)):
    return _compare(payload, db)


class WorkDecisionPayload(WorkComparePayload):
    action: str = Field(min_length=2, max_length=32)
    rationale: str = Field(min_length=5, max_length=1000)


@router.post("/decisions")
def save_decision(payload: WorkDecisionPayload, db: Session = Depends(get_db)):
    calculated = _compare(payload, db)
    if not calculated["feasible"]:
        raise HTTPException(400, "가정 판매량이 가용 재고보다 많습니다. 수량을 수정하세요.")
    product = _product(payload.dataset_id, payload.sku, db)
    rows = db.scalars(select(WorkDailySale).where(WorkDailySale.product_id == product.id,
        WorkDailySale.sale_date.between(payload.as_of_date - timedelta(days=13), payload.as_of_date))).all()
    decision = WorkDecision(product_id=product.id, as_of_date=payload.as_of_date,
                            action=payload.action, rationale=payload.rationale.strip(),
                            observed_units_14d=sum(row.units for row in rows),
                            discount_pct=payload.discount_pct, expected_units_30d=payload.expected_units_30d,
                            purchase_units=payload.purchase_units,
                            assumed_stock=calculated["available_units"] - payload.purchase_units,
                            estimated_gross_profit=calculated["estimated_gross_profit_krw"])
    db.add(decision)
    db.commit()
    return {"id": decision.id, "saved": True, "note": "판단과 계산 당시의 숫자를 저장했습니다. 실제 조치 완료나 성과를 뜻하지 않습니다."}


@router.get("/datasets/{dataset_id}/decisions")
def decisions(dataset_id: int, db: Session = Depends(get_db)):
    _dataset(dataset_id, db)
    rows = db.execute(select(WorkDecision, WorkProduct).join(WorkProduct,
        WorkDecision.product_id == WorkProduct.id).where(WorkProduct.dataset_id == dataset_id)
        .order_by(WorkDecision.created_at.desc(), WorkDecision.id.desc())).all()
    return [{"id": item.id, "sku": product.sku, "name": product.name,
             "as_of_date": item.as_of_date, "action": item.action, "rationale": item.rationale,
             "observed_units_14d": item.observed_units_14d,
             "discount_pct": item.discount_pct, "expected_units_30d": item.expected_units_30d,
             "purchase_units": item.purchase_units, "assumed_stock": item.assumed_stock,
             "estimated_gross_profit_krw": item.estimated_gross_profit,
             "executed_at": item.executed_at, "created_at": item.created_at} for item, product in rows]


class ExecutionPayload(BaseModel):
    executed_at: date


@router.post("/decisions/{decision_id}/execute")
def mark_executed(decision_id: int, payload: ExecutionPayload, db: Session = Depends(get_db)):
    decision = db.get(WorkDecision, decision_id)
    if decision is None:
        raise HTTPException(404, "판단 기록을 찾지 못했습니다.")
    product = db.get(WorkProduct, decision.product_id)
    dataset = db.get(WorkDataset, product.dataset_id)
    if dataset.kind != "company":
        raise HTTPException(400, "H&M 과거 판매에는 실제 조치 실행 기록을 남길 수 없습니다.")
    if payload.executed_at < decision.as_of_date:
        raise HTTPException(400, "실행일은 판단 기준일보다 빠를 수 없습니다.")
    decision.executed_at = payload.executed_at
    db.commit()
    return {"id": decision.id, "executed_at": decision.executed_at}


@router.get("/decisions/{decision_id}/outcome")
def outcome(decision_id: int, db: Session = Depends(get_db)):
    decision = db.get(WorkDecision, decision_id)
    if decision is None:
        raise HTTPException(404, "판단 기록을 찾지 못했습니다.")
    product = db.get(WorkProduct, decision.product_id)
    dataset = db.get(WorkDataset, product.dataset_id)
    if dataset.kind != "company" or decision.executed_at is None:
        return {"status": "not_executed", "note": "자사 자료에서 실제 실행일을 기록한 뒤 비교할 수 있습니다."}
    before_start = decision.executed_at - timedelta(days=7)
    after_end = decision.executed_at + timedelta(days=6)
    if not dataset.complete_calendar or dataset.start_date > before_start or dataset.end_date < after_end:
        return {"status": "waiting_for_data", "note": "실행 전후 7일의 완전한 판매 기록이 필요합니다."}
    rows = db.scalars(select(WorkDailySale).where(WorkDailySale.product_id == product.id,
        WorkDailySale.sale_date.between(before_start, after_end))).all()
    before = sum(row.units for row in rows if row.sale_date < decision.executed_at)
    after = sum(row.units for row in rows if row.sale_date >= decision.executed_at)
    before_rows = [row for row in rows if row.sale_date < decision.executed_at]
    after_rows = [row for row in rows if row.sale_date >= decision.executed_at]
    def observed_amount(key: str, items: list[WorkDailySale]):
        if not items or any(getattr(row, key) is None for row in items):
            return None
        return round(sum(getattr(row, key) for row in items), 2)
    before_snapshot = _input_on(product.id, decision.executed_at - timedelta(days=1), db)
    after_snapshot = _input_on(product.id, after_end, db)
    return {"status": "descriptive", "before_units_7d": before, "after_units_7d": after,
            "before_revenue_7d": observed_amount("revenue", before_rows),
            "after_revenue_7d": observed_amount("revenue", after_rows),
            "before_gross_profit_7d": observed_amount("gross_profit", before_rows),
            "after_gross_profit_7d": observed_amount("gross_profit", after_rows),
            "before_stock": before_snapshot.stock_on_hand if before_snapshot else None,
            "after_stock": after_snapshot.stock_on_hand if after_snapshot else None,
            "before_period": [before_start, decision.executed_at - timedelta(days=1)],
            "after_period": [decision.executed_at, after_end],
            "note": "전후 판매량의 기술적 비교입니다. 조치가 변화를 일으켰다는 인과적 증거는 아닙니다."}
